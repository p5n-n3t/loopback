from __future__ import annotations

import json
import os
import sys
import tempfile
import time
import unittest
from pathlib import Path

RUNTIME = Path(__file__).resolve().parents[1] / "runtime"
sys.path.insert(0, str(RUNTIME))

from audit import AuditStore
from browser import BrowserAdapter
from documents import DocumentTools
from desktop import DesktopAdapter
from dashboard import render_dashboard, render_login
from fleet import FleetStore
from jobs import JobManager, TerminalManager
from policy import ApprovalRequired, PolicyError, PolicyStore


class PolicyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = Path(self.tmp.name)
        self.home = self.cfg / "home"
        self.home.mkdir()
        self.allowed = self.home / "projects"
        self.allowed.mkdir()
        self.denied = self.home / "secret"
        self.denied.mkdir()
        self.policy = PolicyStore(self.cfg)
        self.policy.save({
            "profile": "standard",
            "allowed_roots": [str(self.home)],
            "denied_roots": [str(self.denied)],
            "allow_shell": True,
            "allow_process_control": True,
            "allow_fleet": True,
        })

    def tearDown(self):
        self.tmp.cleanup()

    def test_path_allow_deny(self):
        self.assertEqual(self.policy.check_path(self.allowed), self.allowed.resolve())
        with self.assertRaises(PolicyError):
            self.policy.check_path(self.denied / "x")

    def test_sensitive_command_approval_roundtrip(self):
        with self.assertRaises(ApprovalRequired) as ctx:
            self.policy.check_command("sudo systemctl stop demo")
        approval_id = ctx.exception.approval_id
        self.policy.approve(approval_id)
        self.policy.check_command("sudo systemctl stop demo", approval_id=approval_id)
        with self.assertRaises(ApprovalRequired):
            self.policy.check_command("sudo systemctl stop demo", approval_id=approval_id)

    def test_hard_block_even_trusted(self):
        self.policy.save({"profile": "trusted"})
        with self.assertRaises(PolicyError):
            self.policy.check_command("rm -rf /")


    def test_per_tool_allow_deny(self):
        self.policy.save({"allowed_tools": ["read_file", "diagnostics"], "denied_tools": []})
        self.policy.check_tool("read_file")
        with self.assertRaises(PolicyError):
            self.policy.check_tool("execute")
        self.policy.save({"allowed_tools": ["*"], "denied_tools": ["process_kill"]})
        with self.assertRaises(PolicyError):
            self.policy.check_tool("process_kill")

    def test_browser_disabled_by_default(self):
        with self.assertRaises(PolicyError):
            self.policy.check_browser()
        self.policy.save({"allow_browser": True})
        self.policy.check_browser()

    def test_desktop_disabled_by_default(self):
        with self.assertRaises(PolicyError):
            self.policy.check_desktop()
        self.policy.save({"allow_desktop": True})
        self.policy.check_desktop()


class AuditTests(unittest.TestCase):
    def test_record_and_stats(self):
        with tempfile.TemporaryDirectory() as td:
            a = AuditStore(Path(td))
            a.record("read_file", "ok", duration_ms=2, target="/tmp/x")
            a.record("execute", "error", duration_ms=7)
            rows = a.recent(10)
            self.assertEqual(len(rows), 2)
            stats = a.stats()
            self.assertEqual(stats["total"], 2)
            self.assertEqual(stats["failures"], 1)


class FleetTests(unittest.TestCase):
    def test_node_secrets_are_separate(self):
        with tempfile.TemporaryDirectory() as td:
            f = FleetStore(Path(td))
            f.upsert("ora3", "https://ora3.example/mcp", "secret")
            listed = f.list()
            self.assertEqual(listed[0]["name"], "ora3")
            self.assertNotIn("token", listed[0])
            node, token = f.get("ora3")
            self.assertEqual(token, "secret")
            self.assertTrue(f.remove("ora3"))


class JobTests(unittest.TestCase):
    def test_background_job_output(self):
        with tempfile.TemporaryDirectory() as td:
            j = JobManager(Path(td), "/bin/sh")
            info = j.start("printf hello", td)
            for _ in range(30):
                current = j.info(info["id"])
                if not current["running"]:
                    break
                time.sleep(0.05)
            out = j.output(info["id"])
            self.assertIn("hello", out["content"])


@unittest.skipUnless(sys.platform.startswith("linux"), "PTY test requires Unix")
class TerminalTests(unittest.TestCase):
    def test_terminal_roundtrip(self):
        t = TerminalManager("/bin/sh")
        with tempfile.TemporaryDirectory() as td:
            info = t.start(td, shell="/bin/sh")
            try:
                t.read(info["id"], clear=True)
                t.write(info["id"], "printf LOOPBACK_PTY_OK\\n")
                data = ""
                for _ in range(30):
                    time.sleep(0.05)
                    data += t.read(info["id"])["content"]
                    if "LOOPBACK_PTY_OK" in data:
                        break
                self.assertIn("LOOPBACK_PTY_OK", data)
            finally:
                t.close(info["id"])


class DocumentTests(unittest.TestCase):
    def test_docx_xlsx_pdf_roundtrip(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter

        tools = DocumentTools()
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)

            docx = root / "sample.docx"
            doc = Document()
            doc.add_paragraph("hello Loopback")
            doc.save(docx)
            self.assertIn("hello Loopback", tools.docx_text(docx)["paragraphs"])
            replaced = root / "replaced.docx"
            result = tools.docx_replace_text(docx, "Loopback", "world", replaced)
            self.assertEqual(result["replacements"], 1)
            self.assertIn("hello world", tools.docx_text(replaced)["paragraphs"])

            xlsx = root / "sample.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            ws["A1"] = "one"
            wb.save(xlsx)
            self.assertEqual(tools.xlsx_read_range(xlsx, "Data", "A1:A1")["values"], [["one"]])
            out_xlsx = root / "out.xlsx"
            tools.xlsx_write_range(xlsx, "Data", "B2", [[1, 2], [3, 4]], out_xlsx)
            self.assertEqual(
                tools.xlsx_read_range(out_xlsx, "Data", "B2:C3")["values"],
                [[1, 2], [3, 4]],
            )

            pdf = root / "sample.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.add_blank_page(width=100, height=100)
            with pdf.open("wb") as handle:
                writer.write(handle)
            extracted = root / "one.pdf"
            result = tools.pdf_extract_pages(pdf, [2], extracted)
            self.assertEqual(result["pages"], [2])
            self.assertEqual(tools.pdf_text(extracted)["total_pages"], 1)


class BrowserAdapterTests(unittest.TestCase):
    def test_status_and_validation(self):
        browser = BrowserAdapter("definitely-not-installed-loopback-browser")
        self.assertFalse(browser.status()["available"])
        with self.assertRaises(ValueError):
            browser.open("http://example.com")


class DesktopAdapterTests(unittest.TestCase):
    def test_status_shape(self):
        status = DesktopAdapter().status()
        self.assertIn("available", status)
        self.assertIn("platform", status)
        self.assertIn("xdotool", status)

class DashboardTests(unittest.TestCase):
    def test_render(self):
        self.assertIn("LOOPBACK CONTROL PLANE", render_login("example.test"))
        html = render_dashboard("example.test")
        self.assertIn("Recent MCP activity", html)
        self.assertIn("Persistent terminals", html)
        self.assertIn("Nodes", html)


if __name__ == "__main__":
    unittest.main()
