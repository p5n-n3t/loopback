from __future__ import annotations

from pathlib import Path
from typing import Any



def _docx():
    from docx import Document  # lazy: keeps idle RSS low until a document tool is used

    return Document


def _openpyxl():
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    return load_workbook, range_boundaries


def _pypdf():
    from pypdf import PdfReader, PdfWriter

    return PdfReader, PdfWriter


class DocumentTools:
    """DOCX/XLSX/PDF helpers. Third-party libraries are imported on first use."""
    def docx_text(self, path: Path) -> dict[str, Any]:
        doc = _docx()(str(path))
        paragraphs = [p.text for p in doc.paragraphs]
        tables = []
        for table in doc.tables:
            tables.append([[cell.text for cell in row.cells] for row in table.rows])
        return {
            "path": str(path),
            "paragraphs": paragraphs,
            "tables": tables,
        }

    def docx_replace_text(
        self,
        path: Path,
        old_text: str,
        new_text: str,
        output: Path,
    ) -> dict[str, Any]:
        if not old_text:
            raise ValueError("old_text cannot be empty")
        doc = _docx()(str(path))
        replacements = 0

        def replace_paragraph(paragraph) -> None:
            nonlocal replacements
            full = "".join(run.text for run in paragraph.runs)
            count = full.count(old_text)
            if count == 0:
                return
            replaced = full.replace(old_text, new_text)
            if paragraph.runs:
                paragraph.runs[0].text = replaced
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = replaced
            replacements += count

        for paragraph in doc.paragraphs:
            replace_paragraph(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_paragraph(paragraph)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output))
        return {
            "input": str(path),
            "output": str(output),
            "replacements": replacements,
            "formatting_note": "replacement spanning multiple runs is consolidated into the first run",
        }

    def xlsx_read_range(self, path: Path, sheet: str, cell_range: str) -> dict[str, Any]:
        wb = _openpyxl()[0](str(path), data_only=False, read_only=True)
        if sheet not in wb.sheetnames:
            raise KeyError(f"sheet not found: {sheet}")
        ws = wb[sheet]
        min_col, min_row, max_col, max_row = _openpyxl()[1](cell_range)
        values = [
            [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
            for r in range(min_row, max_row + 1)
        ]
        return {
            "path": str(path),
            "sheet": sheet,
            "range": cell_range,
            "values": values,
        }

    def xlsx_write_range(
        self,
        path: Path,
        sheet: str,
        start_cell: str,
        values: list[list[Any]],
        output: Path,
    ) -> dict[str, Any]:
        wb = _openpyxl()[0](str(path))
        if sheet not in wb.sheetnames:
            raise KeyError(f"sheet not found: {sheet}")
        ws = wb[sheet]
        min_col, min_row, _, _ = _openpyxl()[1](f"{start_cell}:{start_cell}")
        rows = 0
        cells = 0
        for r_off, row in enumerate(values):
            rows += 1
            for c_off, value in enumerate(row):
                ws.cell(row=min_row + r_off, column=min_col + c_off).value = value
                cells += 1
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output))
        return {
            "input": str(path),
            "output": str(output),
            "sheet": sheet,
            "start_cell": start_cell,
            "rows": rows,
            "cells": cells,
        }

    def pdf_text(self, path: Path, start_page: int = 1, max_pages: int = 20) -> dict[str, Any]:
        reader = _pypdf()[0](str(path))
        start = max(1, int(start_page))
        count = max(1, min(int(max_pages), 200))
        pages = []
        for index in range(start - 1, min(len(reader.pages), start - 1 + count)):
            pages.append({
                "page": index + 1,
                "text": reader.pages[index].extract_text() or "",
            })
        return {
            "path": str(path),
            "total_pages": len(reader.pages),
            "pages": pages,
        }

    def pdf_merge(self, paths: list[Path], output: Path) -> dict[str, Any]:
        if not paths:
            raise ValueError("at least one PDF is required")
        writer = _pypdf()[1]()
        total = 0
        for path in paths:
            reader = _pypdf()[0](str(path))
            for page in reader.pages:
                writer.add_page(page)
                total += 1
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("wb") as handle:
            writer.write(handle)
        return {"inputs": [str(p) for p in paths], "output": str(output), "pages": total}

    def pdf_extract_pages(self, path: Path, pages: list[int], output: Path) -> dict[str, Any]:
        reader = _pypdf()[0](str(path))
        writer = _pypdf()[1]()
        selected = []
        for page_number in pages:
            n = int(page_number)
            if n < 1 or n > len(reader.pages):
                raise ValueError(f"page out of range: {n}")
            writer.add_page(reader.pages[n - 1])
            selected.append(n)
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("wb") as handle:
            writer.write(handle)
        return {"input": str(path), "output": str(output), "pages": selected}
